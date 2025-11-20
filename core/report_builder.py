"""Report builder for generating custom reports."""

from __future__ import annotations

import json
import logging
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)


class ReportFormat(Enum):
    """Report export formats."""
    PDF = "pdf"
    HTML = "html"
    DOCX = "docx"
    CSV = "csv"
    EXCEL = "xlsx"


@dataclass
class ReportSection:
    """A section in a report."""
    title: str
    content_type: str  # "text", "table", "chart", "stats"
    content: dict[str, Any] = field(default_factory=dict)
    order: int = 0


@dataclass
class ReportTemplate:
    """A report template definition."""
    name: str
    description: str
    sections: list[ReportSection] = field(default_factory=list)
    filters: dict[str, Any] = field(default_factory=dict)


class ReportBuilder:
    """Builds custom reports from templates and data."""
    
    def __init__(self, session_manager):
        """Initialize report builder."""
        self.session_manager = session_manager
        self.templates_dir = Path("data/report_templates")
        self.templates_dir.mkdir(parents=True, exist_ok=True)
    
    def get_available_templates(self) -> list[ReportTemplate]:
        """Get list of available report templates."""
        templates = []
        
        # Session Summary Template
        templates.append(ReportTemplate(
            name="Session Summary",
            description="Summary report for a single session with key metrics and shot data",
            sections=[
                ReportSection("Session Information", "text", order=1),
                ReportSection("Key Statistics", "stats", order=2),
                ReportSection("Shot Data", "table", order=3),
                ReportSection("Performance Trends", "chart", order=4),
            ]
        ))
        
        # Performance Trends Template
        templates.append(ReportTemplate(
            name="Performance Trends",
            description="Trend analysis across multiple sessions with charts",
            sections=[
                ReportSection("Summary Statistics", "stats", order=1),
                ReportSection("Speed Trends", "chart", order=2),
                ReportSection("Distance Trends", "chart", order=3),
                ReportSection("Spin Rate Trends", "chart", order=4),
                ReportSection("Dispersion Analysis", "chart", order=5),
            ]
        ))
        
        # Club Comparison Template
        templates.append(ReportTemplate(
            name="Club Comparison",
            description="Compare performance across different clubs",
            sections=[
                ReportSection("Club Statistics", "stats", order=1),
                ReportSection("Club Comparison Chart", "chart", order=2),
                ReportSection("Club Performance Table", "table", order=3),
            ]
        ))
        
        # Consistency Analysis Template
        templates.append(ReportTemplate(
            name="Consistency Analysis",
            description="Analyze shot consistency with standard deviation and variance metrics",
            sections=[
                ReportSection("Consistency Overview", "stats", order=1),
                ReportSection("Distribution Analysis", "chart", order=2),
                ReportSection("Box Plot Analysis", "chart", order=3),
                ReportSection("Consistency Metrics Table", "table", order=4),
            ]
        ))
        
        # Dispersion Analysis Template
        templates.append(ReportTemplate(
            name="Dispersion Analysis",
            description="Analyze shot dispersion patterns with heat maps and scatter plots",
            sections=[
                ReportSection("Dispersion Overview", "stats", order=1),
                ReportSection("Dispersion Heat Map", "chart", order=2),
                ReportSection("Dispersion Statistics", "table", order=3),
            ]
        ))
        
        # Goal Progress Template
        templates.append(ReportTemplate(
            name="Goal Progress",
            description="Track progress toward goals with statistics and trends",
            sections=[
                ReportSection("Goal Summary", "text", order=1),
                ReportSection("Progress Statistics", "stats", order=2),
                ReportSection("Progress Trends", "chart", order=3),
                ReportSection("Goal Details", "table", order=4),
            ]
        ))
        
        # Shot Comparison Template
        templates.append(ReportTemplate(
            name="Shot Comparison",
            description="Compare multiple shots side-by-side with key metrics",
            sections=[
                ReportSection("Comparison Overview", "stats", order=1),
                ReportSection("Metric Comparison Chart", "chart", order=2),
                ReportSection("Detailed Comparison Table", "table", order=3),
            ]
        ))
        
        # Practice Session Template
        templates.append(ReportTemplate(
            name="Practice Session",
            description="Summary of a practice session with key insights",
            sections=[
                ReportSection("Session Overview", "text", order=1),
                ReportSection("Session Statistics", "stats", order=2),
                ReportSection("Performance Trends", "chart", order=3),
                ReportSection("Shot List", "table", order=4),
            ]
        ))
        
        # Advanced Analytics Template
        templates.append(ReportTemplate(
            name="Advanced Analytics",
            description="Comprehensive analytics including tempo, swing plane, and ball flight prediction",
            sections=[
                ReportSection("Analytics Summary", "stats", order=1),
                ReportSection("Tempo Analysis", "text", order=2),
                ReportSection("Swing Plane Analysis", "text", order=3),
                ReportSection("Ball Flight Prediction", "text", order=4),
            ]
        ))
        
        # Shot Dispersion Template
        templates.append(ReportTemplate(
            name="Shot Dispersion",
            description="Detailed dispersion analysis with heat maps and pattern recognition",
            sections=[
                ReportSection("Dispersion Overview", "stats", order=1),
                ReportSection("Dispersion Heat Map", "chart", order=2),
                ReportSection("Dispersion Statistics", "table", order=3),
            ]
        ))
        
        # Custom Template (user-defined)
        templates.append(ReportTemplate(
            name="Custom Report",
            description="Build your own custom report",
            sections=[]
        ))
        
        return templates
    
    def generate_report(
        self,
        template: ReportTemplate,
        output_path: Path,
        format: ReportFormat = ReportFormat.PDF,
        session_ids: Optional[list[int]] = None,
        shot_ids: Optional[list[int]] = None,
        date_range: Optional[tuple[datetime, datetime]] = None,
        club_filter: Optional[str] = None,
        include_charts: bool = True,
        include_stats: bool = True,
        include_tables: bool = True,
    ) -> bool:
        """Generate a report from a template."""
        try:
            if format == ReportFormat.PDF:
                return self._generate_pdf(template, output_path, session_ids, shot_ids, date_range, club_filter, include_charts, include_stats, include_tables)
            elif format == ReportFormat.HTML:
                return self._generate_html(template, output_path, session_ids, shot_ids, date_range, club_filter, include_charts, include_stats, include_tables)
            elif format == ReportFormat.DOCX:
                return self._generate_docx(template, output_path, session_ids, shot_ids, date_range, club_filter, include_charts, include_stats, include_tables)
            elif format == ReportFormat.CSV:
                return self._generate_csv(template, output_path, session_ids, shot_ids, date_range, club_filter)
            elif format == ReportFormat.EXCEL:
                return self._generate_excel(template, output_path, session_ids, shot_ids, date_range, club_filter)
            else:
                logger.error("Unsupported report format: %s", format)
                return False
        except Exception as e:
            logger.error("Error generating report: %s", e, exc_info=True)
            return False
    
    def _generate_pdf(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
        include_charts: bool = True,
        include_stats: bool = True,
        include_tables: bool = True,
    ) -> bool:
        """Generate PDF report."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image
            from reportlab.pdfgen import canvas
            
            doc = SimpleDocTemplate(str(output_path), pagesize=letter)
            story = []
            styles = getSampleStyleSheet()
            
            # Title style
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#ff4d4d'),
                spaceAfter=30,
            )
            
            # Add title
            story.append(Paragraph(f"{template.name} Report", title_style))
            story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
            story.append(Spacer(1, 0.3*inch))
            
            # Get data
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            # Generate sections
            for section in sorted(template.sections, key=lambda s: s.order):
                if section.content_type == "text":
                    story.append(Paragraph(f"<b>{section.title}</b>", styles['Heading2']))
                    story.append(self._generate_text_section(section, sessions, shots, styles))
                elif section.content_type == "stats" and include_stats:
                    story.append(Paragraph(f"<b>{section.title}</b>", styles['Heading2']))
                    story.append(self._generate_stats_section(section, shots, styles))
                elif section.content_type == "table" and include_tables:
                    story.append(Paragraph(f"<b>{section.title}</b>", styles['Heading2']))
                    story.append(self._generate_table_section(section, shots, styles))
                elif section.content_type == "chart" and include_charts:
                    story.append(Paragraph(f"<b>{section.title}</b>", styles['Heading2']))
                    chart_path = self._generate_chart_image(section, shots, sessions)
                    if chart_path and chart_path.exists():
                        from reportlab.platypus import Image
                        img = Image(str(chart_path), width=6*inch, height=4*inch)
                        story.append(img)
                    else:
                        story.append(Paragraph("Chart could not be generated.", styles['Normal']))
                
                story.append(Spacer(1, 0.2*inch))
            
            doc.build(story)
            logger.info("Generated PDF report: %s", output_path)
            return True
        except ImportError:
            logger.error("reportlab not installed. Install with: pip install reportlab")
            return False
        except Exception as e:
            logger.error("Error generating PDF report: %s", e, exc_info=True)
            return False
    
    def _generate_html(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
        include_charts: bool = True,
        include_stats: bool = True,
        include_tables: bool = True,
    ) -> bool:
        """Generate HTML report."""
        try:
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{template.name} Report</title>
    <style>
        body {{
            font-family: Arial, sans-serif;
            margin: 40px;
            background-color: #0f1116;
            color: #c9d1d9;
        }}
        h1 {{
            color: #ff4d4d;
            border-bottom: 2px solid #30363d;
            padding-bottom: 10px;
        }}
        h2 {{
            color: #c9d1d9;
            margin-top: 30px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
        }}
        th, td {{
            border: 1px solid #30363d;
            padding: 12px;
            text-align: left;
        }}
        th {{
            background-color: #21262d;
            color: #ff4d4d;
            font-weight: bold;
        }}
        tr:nth-child(even) {{
            background-color: #161b22;
        }}
        .stats {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 20px;
            margin: 20px 0;
        }}
        .stat-item {{
            background-color: #21262d;
            padding: 15px;
            border-radius: 8px;
            border: 1px solid #30363d;
        }}
        .stat-label {{
            color: #8b949e;
            font-size: 12px;
            margin-bottom: 5px;
        }}
        .stat-value {{
            color: #c9d1d9;
            font-size: 24px;
            font-weight: bold;
        }}
    </style>
</head>
<body>
    <h1>{template.name} Report</h1>
    <p>Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
"""
            
            # Generate sections
            for section in sorted(template.sections, key=lambda s: s.order):
                html += f"<h2>{section.title}</h2>\n"
                
                if section.content_type == "text":
                    html += self._generate_text_section_html(section, sessions, shots)
                elif section.content_type == "stats" and include_stats:
                    html += self._generate_stats_section_html(section, shots)
                elif section.content_type == "table" and include_tables:
                    html += self._generate_table_section_html(section, shots)
                elif section.content_type == "chart" and include_charts:
                    chart_path = self._generate_chart_image(section, shots, sessions)
                    if chart_path and chart_path.exists():
                        import base64
                        with open(chart_path, 'rb') as f:
                            img_data = base64.b64encode(f.read()).decode('utf-8')
                        html += f'<img src="data:image/png;base64,{img_data}" style="max-width: 100%; height: auto;" alt="{section.title}">\n'
                    else:
                        html += "<p>Chart could not be generated.</p>\n"
            
            html += """
</body>
</html>
"""
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(html)
            
            logger.info("Generated HTML report: %s", output_path)
            return True
        except Exception as e:
            logger.error("Error generating HTML report: %s", e, exc_info=True)
            return False
    
    def _generate_docx(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
        include_charts: bool = True,
        include_stats: bool = True,
        include_tables: bool = True,
    ) -> bool:
        """Generate DOCX report."""
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            
            doc = Document()
            
            # Title
            title = doc.add_heading(template.name + " Report", 0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            title.runs[0].font.color.rgb = RGBColor(255, 77, 77)
            
            # Date
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph()
            
            # Get data
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            # Generate sections
            for section in sorted(template.sections, key=lambda s: s.order):
                doc.add_heading(section.title, level=1)
                
                if section.content_type == "text":
                    self._generate_text_section_docx(doc, section, sessions, shots)
                elif section.content_type == "stats" and include_stats:
                    self._generate_stats_section_docx(doc, section, shots)
                elif section.content_type == "table" and include_tables:
                    self._generate_table_section_docx(doc, section, shots)
                elif section.content_type == "chart" and include_charts:
                    chart_path = self._generate_chart_image(section, shots, sessions)
                    if chart_path and chart_path.exists():
                        doc.add_picture(str(chart_path), width=Inches(6))
                    else:
                        doc.add_paragraph("Chart could not be generated.")
                
                doc.add_paragraph()
            
            doc.save(str(output_path))
            logger.info("Generated DOCX report: %s", output_path)
            return True
        except ImportError:
            logger.error("python-docx not installed. Install with: pip install python-docx")
            return False
        except Exception as e:
            logger.error("Error generating DOCX report: %s", e, exc_info=True)
            return False
    
    def _generate_csv(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
    ) -> bool:
        """Generate CSV report."""
        try:
            import csv
            from sqlalchemy.orm import Session
            from core.session_manager import ShotModel, SessionModel
            
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            if not shots:
                logger.warning("No shots to export to CSV")
                return False
            
            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                
                # Write header
                headers = [
                    "Date", "Time", "Session", "Club", "Club Speed (mph)", "Ball Speed (mph)",
                    "Launch Angle (deg)", "Spin Rate (rpm)", "Carry Distance (yds)",
                    "Total Distance (yds)", "Side Spin (rpm)", "Back Spin (rpm)",
                    "Launch Direction (deg)", "Apex Height (yds)", "Descent Angle (deg)",
                    "Smash Factor", "Dynamic Loft (deg)", "Attack Angle (deg)",
                    "Club Path (deg)", "Face Angle (deg)", "Notes"
                ]
                writer.writerow(headers)
                
                # Get session map
                session_map = {}
                club_map = {}
                with Session(self.session_manager.engine) as session:
                    for shot in shots:
                        if shot.session_id not in session_map:
                            sess = session.get(SessionModel, shot.session_id)
                            session_map[shot.session_id] = sess.name if sess else "Unknown"
                            club_map[shot.session_id] = sess.club if sess and sess.club else ""
                
                # Write data rows
                for shot in shots:
                    row = [
                        shot.recorded_at.strftime("%Y-%m-%d"),
                        shot.recorded_at.strftime("%H:%M:%S"),
                        session_map.get(shot.session_id, "Unknown"),
                        club_map.get(shot.session_id, ""),
                        f"{shot.club_speed:.1f}" if shot.club_speed else "",
                        f"{shot.ball_speed:.1f}" if shot.ball_speed else "",
                        f"{shot.launch_angle:.1f}" if shot.launch_angle else "",
                        f"{shot.spin_rate:.0f}" if shot.spin_rate else "",
                        f"{shot.carry_distance:.1f}" if shot.carry_distance else "",
                        f"{shot.total_distance:.1f}" if shot.total_distance else "",
                        f"{shot.side_spin:.0f}" if shot.side_spin else "",
                        f"{shot.back_spin:.0f}" if shot.back_spin else "",
                        f"{shot.launch_direction:.1f}" if shot.launch_direction else "",
                        f"{shot.apex_height:.1f}" if shot.apex_height else "",
                        f"{shot.descent_angle:.1f}" if shot.descent_angle else "",
                        f"{shot.smash_factor:.2f}" if shot.smash_factor else "",
                        f"{shot.dynamic_loft:.1f}" if shot.dynamic_loft else "",
                        f"{shot.attack_angle:.1f}" if shot.attack_angle else "",
                        f"{shot.club_path:.1f}" if shot.club_path else "",
                        f"{shot.face_angle:.1f}" if shot.face_angle else "",
                        shot.notes or "",
                    ]
                    writer.writerow(row)
            
            logger.info("Generated CSV report: %s", output_path)
            return True
        except Exception as e:
            logger.error("Error generating CSV report: %s", e, exc_info=True)
            return False
    
    def _generate_excel(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
    ) -> bool:
        """Generate Excel report."""
        try:
            import pandas as pd
            from sqlalchemy.orm import Session
            from core.session_manager import ShotModel, SessionModel
            
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            if not shots:
                logger.warning("No shots to export to Excel")
                return False
            
            # Get session map
            session_map = {}
            club_map = {}
            with Session(self.session_manager.engine) as session:
                for shot in shots:
                    if shot.session_id not in session_map:
                        sess = session.get(SessionModel, shot.session_id)
                        session_map[shot.session_id] = sess.name if sess else "Unknown"
                        club_map[shot.session_id] = sess.club if sess and sess.club else ""
            
            # Build data rows
            rows = []
            for shot in shots:
                row = {
                    "Date": shot.recorded_at.strftime("%Y-%m-%d"),
                    "Time": shot.recorded_at.strftime("%H:%M:%S"),
                    "Session": session_map.get(shot.session_id, "Unknown"),
                    "Club": club_map.get(shot.session_id, ""),
                    "Club Speed (mph)": shot.club_speed,
                    "Ball Speed (mph)": shot.ball_speed,
                    "Launch Angle (deg)": shot.launch_angle,
                    "Spin Rate (rpm)": shot.spin_rate,
                    "Carry Distance (yds)": shot.carry_distance,
                    "Total Distance (yds)": shot.total_distance,
                    "Side Spin (rpm)": shot.side_spin,
                    "Back Spin (rpm)": shot.back_spin,
                    "Launch Direction (deg)": shot.launch_direction,
                    "Apex Height (yds)": shot.apex_height,
                    "Descent Angle (deg)": shot.descent_angle,
                    "Smash Factor": shot.smash_factor,
                    "Dynamic Loft (deg)": shot.dynamic_loft,
                    "Attack Angle (deg)": shot.attack_angle,
                    "Club Path (deg)": shot.club_path,
                    "Face Angle (deg)": shot.face_angle,
                    "Notes": shot.notes or "",
                }
                rows.append(row)
            
            df = pd.DataFrame(rows)
            df.to_excel(output_path, index=False, engine='openpyxl')
            
            logger.info("Generated Excel report: %s", output_path)
            return True
        except ImportError:
            logger.error("pandas and openpyxl required for Excel export. Install with: pip install pandas openpyxl")
            return False
        except Exception as e:
            logger.error("Error generating Excel report: %s", e, exc_info=True)
            return False
    
    def _get_report_data(
        self,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
    ) -> tuple[list, list]:
        """Get shots and sessions for report."""
        from sqlalchemy.orm import Session
        from core.session_manager import SessionModel, ShotModel
        
        with Session(self.session_manager.engine) as session:
            query = session.query(ShotModel)
            
            if shot_ids:
                query = query.filter(ShotModel.id.in_(shot_ids))
            elif session_ids:
                query = query.filter(ShotModel.session_id.in_(session_ids))
            
            if date_range:
                start, end = date_range
                query = query.join(SessionModel).filter(
                    SessionModel.started_at >= start,
                    SessionModel.started_at <= end
                )
            
            if club_filter:
                query = query.join(SessionModel).filter(SessionModel.club == club_filter)
            
            shots = query.all()
            
            # Get unique sessions
            session_ids_found = list(set(s.session_id for s in shots))
            sessions = session.query(SessionModel).filter(SessionModel.id.in_(session_ids_found)).all() if session_ids_found else []
            
            return shots, sessions
    
    def _generate_text_section(self, section: ReportSection, sessions: list, shots: list, styles) -> Paragraph:
        """Generate text section content."""
        if section.title == "Session Information" and sessions:
            sess = sessions[0]
            text = f"""
            <b>Name:</b> {sess.name}<br/>
            <b>Club:</b> {sess.club or 'N/A'}<br/>
            <b>Started:</b> {sess.started_at.strftime('%Y-%m-%d %H:%M:%S')}<br/>
            <b>Shots:</b> {len(shots)}<br/>
            """
            if sess.notes:
                text += f"<b>Notes:</b> {sess.notes}<br/>"
            return Paragraph(text, styles['Normal'])
        return Paragraph("No data available.", styles['Normal'])
    
    def _generate_stats_section(self, section: ReportSection, shots: list, styles) -> Table:
        """Generate statistics section."""
        if not shots:
            return Paragraph("No data available.", styles['Normal'])
        
        club_speeds = [s.club_speed for s in shots if s.club_speed]
        ball_speeds = [s.ball_speed for s in shots if s.ball_speed]
        carries = [s.carry_distance for s in shots if s.carry_distance]
        
        stats_data = [
            ["Metric", "Value"],
            ["Total Shots", str(len(shots))],
            ["Avg Club Speed", f"{sum(club_speeds)/len(club_speeds):.1f} mph" if club_speeds else "N/A"],
            ["Avg Ball Speed", f"{sum(ball_speeds)/len(ball_speeds):.1f} mph" if ball_speeds else "N/A"],
            ["Avg Carry", f"{sum(carries)/len(carries):.1f} yds" if carries else "N/A"],
            ["Max Carry", f"{max(carries):.1f} yds" if carries else "N/A"],
        ]
        
        table = Table(stats_data, colWidths=[2*inch, 2*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), '#21262d'),
            ('TEXTCOLOR', (0, 0), (-1, 0), '#ff4d4d'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), '#161b22'),
            ('TEXTCOLOR', (0, 1), (-1, -1), '#c9d1d9'),
            ('GRID', (0, 0), (-1, -1), 1, '#30363d'),
        ]))
        return table
    
    def _generate_table_section(self, section: ReportSection, shots: list, styles) -> Table:
        """Generate table section."""
        if not shots:
            return Paragraph("No data available.", styles['Normal'])
        
        headers = ["Date", "Club Speed", "Ball Speed", "Carry", "Total", "Spin Rate"]
        data = [headers]
        
        for shot in shots[:50]:  # Limit to 50 shots for PDF
            row = [
                shot.recorded_at.strftime("%Y-%m-%d %H:%M"),
                f"{shot.club_speed:.1f}" if shot.club_speed else "N/A",
                f"{shot.ball_speed:.1f}" if shot.ball_speed else "N/A",
                f"{shot.carry_distance:.1f}" if shot.carry_distance else "N/A",
                f"{shot.total_distance:.1f}" if shot.total_distance else "N/A",
                f"{shot.spin_rate:.0f}" if shot.spin_rate else "N/A",
            ]
            data.append(row)
        
        table = Table(data, colWidths=[1.2*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), '#21262d'),
            ('TEXTCOLOR', (0, 0), (-1, 0), '#ff4d4d'),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), '#161b22'),
            ('TEXTCOLOR', (0, 1), (-1, -1), '#c9d1d9'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('GRID', (0, 0), (-1, -1), 1, '#30363d'),
        ]))
        return table
    
    def _generate_text_section_html(self, section: ReportSection, sessions: list, shots: list) -> str:
        """Generate HTML text section."""
        if section.title == "Session Information" and sessions:
            sess = sessions[0]
            return f"""
            <p><b>Name:</b> {sess.name}</p>
            <p><b>Club:</b> {sess.club or 'N/A'}</p>
            <p><b>Started:</b> {sess.started_at.strftime('%Y-%m-%d %H:%M:%S')}</p>
            <p><b>Shots:</b> {len(shots)}</p>
            {f'<p><b>Notes:</b> {sess.notes}</p>' if sess.notes else ''}
            """
        return "<p>No data available.</p>"
    
    def _generate_stats_section_html(self, section: ReportSection, shots: list) -> str:
        """Generate HTML stats section."""
        if not shots:
            return "<p>No data available.</p>"
        
        club_speeds = [s.club_speed for s in shots if s.club_speed]
        ball_speeds = [s.ball_speed for s in shots if s.ball_speed]
        carries = [s.carry_distance for s in shots if s.carry_distance]
        
        return f"""
        <div class="stats">
            <div class="stat-item">
                <div class="stat-label">Total Shots</div>
                <div class="stat-value">{len(shots)}</div>
            </div>
            <div class="stat-item">
                <div class="stat-label">Avg Club Speed</div>
                <div class="stat-value">{sum(club_speeds)/len(club_speeds):.1f} mph</div>
            </div>
            <div class="stat-item">
                <div class="stat-label">Avg Ball Speed</div>
                <div class="stat-value">{sum(ball_speeds)/len(ball_speeds):.1f} mph</div>
            </div>
            <div class="stat-item">
                <div class="stat-label">Avg Carry</div>
                <div class="stat-value">{sum(carries)/len(carries):.1f} yds</div>
            </div>
            <div class="stat-item">
                <div class="stat-label">Max Carry</div>
                <div class="stat-value">{max(carries):.1f} yds</div>
            </div>
        </div>
        """
    
    def _generate_table_section_html(self, section: ReportSection, shots: list) -> str:
        """Generate HTML table section."""
        if not shots:
            return "<p>No data available.</p>"
        
        html = "<table><thead><tr><th>Date</th><th>Club Speed</th><th>Ball Speed</th><th>Carry</th><th>Total</th><th>Spin Rate</th></tr></thead><tbody>"
        
        for shot in shots[:100]:  # Limit to 100 shots
            html += f"""
            <tr>
                <td>{shot.recorded_at.strftime("%Y-%m-%d %H:%M")}</td>
                <td>{f'{shot.club_speed:.1f}' if shot.club_speed else 'N/A'}</td>
                <td>{f'{shot.ball_speed:.1f}' if shot.ball_speed else 'N/A'}</td>
                <td>{f'{shot.carry_distance:.1f}' if shot.carry_distance else 'N/A'}</td>
                <td>{f'{shot.total_distance:.1f}' if shot.total_distance else 'N/A'}</td>
                <td>{f'{shot.spin_rate:.0f}' if shot.spin_rate else 'N/A'}</td>
            </tr>
            """
        
        html += "</tbody></table>"
        return html
    
    def _generate_text_section_docx(self, doc, section: ReportSection, sessions: list, shots: list) -> None:
        """Generate DOCX text section."""
        if section.title == "Session Information" and sessions:
            sess = sessions[0]
            doc.add_paragraph(f"Name: {sess.name}")
            doc.add_paragraph(f"Club: {sess.club or 'N/A'}")
            doc.add_paragraph(f"Started: {sess.started_at.strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph(f"Shots: {len(shots)}")
            if sess.notes:
                doc.add_paragraph(f"Notes: {sess.notes}")
    
    def _generate_stats_section_docx(self, doc, section: ReportSection, shots: list) -> None:
        """Generate DOCX stats section."""
        if not shots:
            doc.add_paragraph("No data available.")
            return
        
        club_speeds = [s.club_speed for s in shots if s.club_speed]
        ball_speeds = [s.ball_speed for s in shots if s.ball_speed]
        carries = [s.carry_distance for s in shots if s.carry_distance]
        
        doc.add_paragraph(f"Total Shots: {len(shots)}")
        doc.add_paragraph(f"Avg Club Speed: {sum(club_speeds)/len(club_speeds):.1f} mph" if club_speeds else "N/A")
        doc.add_paragraph(f"Avg Ball Speed: {sum(ball_speeds)/len(ball_speeds):.1f} mph" if ball_speeds else "N/A")
        doc.add_paragraph(f"Avg Carry: {sum(carries)/len(carries):.1f} yds" if carries else "N/A")
        doc.add_paragraph(f"Max Carry: {max(carries):.1f} yds" if carries else "N/A")
    
    def _generate_table_section_docx(self, doc, section: ReportSection, shots: list) -> None:
        """Generate DOCX table section."""
        if not shots:
            doc.add_paragraph("No data available.")
            return
        
        table = doc.add_table(rows=1, cols=6)
        table.style = 'Light Grid Accent 1'
        
        # Header row
        header_cells = table.rows[0].cells
        header_cells[0].text = "Date"
        header_cells[1].text = "Club Speed"
        header_cells[2].text = "Ball Speed"
        header_cells[3].text = "Carry"
        header_cells[4].text = "Total"
        header_cells[5].text = "Spin Rate"
        
        # Data rows
        for shot in shots[:100]:  # Limit to 100 shots
            row_cells = table.add_row().cells
            row_cells[0].text = shot.recorded_at.strftime("%Y-%m-%d %H:%M")
            row_cells[1].text = f"{shot.club_speed:.1f}" if shot.club_speed else "N/A"
            row_cells[2].text = f"{shot.ball_speed:.1f}" if shot.ball_speed else "N/A"
            row_cells[3].text = f"{shot.carry_distance:.1f}" if shot.carry_distance else "N/A"
            row_cells[4].text = f"{shot.total_distance:.1f}" if shot.total_distance else "N/A"
            row_cells[5].text = f"{shot.spin_rate:.0f}" if shot.spin_rate else "N/A"
    
    def _generate_chart_image(self, section: ReportSection, shots: list, sessions: list) -> Optional[Path]:
        """Generate chart image for a section.
        
        Args:
            section: Report section with chart type
            shots: List of shot data
            sessions: List of session data
            
        Returns:
            Path to generated chart image, or None if generation failed
        """
        try:
            import matplotlib
            matplotlib.use('Agg')  # Non-interactive backend
            import matplotlib.pyplot as plt
            import numpy as np
        except ImportError:
            logger.warning("matplotlib not available for chart generation")
            return None
        
        if not shots:
            return None
        
        # Create temp directory for charts
        temp_dir = Path("data/temp_charts")
        temp_dir.mkdir(parents=True, exist_ok=True)
        
        # Generate chart based on section title
        chart_type = section.title.lower()
        chart_path = temp_dir / f"chart_{section.title.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
        
        try:
            fig, ax = plt.subplots(figsize=(10, 6))
            fig.patch.set_facecolor('#0f1116')
            ax.set_facecolor('#161b22')
            ax.tick_params(colors='#c9d1d9')
            ax.spines['bottom'].set_color('#30363d')
            ax.spines['top'].set_color('#30363d')
            ax.spines['right'].set_color('#30363d')
            ax.spines['left'].set_color('#30363d')
            ax.xaxis.label.set_color('#c9d1d9')
            ax.yaxis.label.set_color('#c9d1d9')
            ax.title.set_color('#c9d1d9')
            
            if "performance trends" in chart_type or "trend" in chart_type:
                # Line chart of carry distance over time
                dates = [s.recorded_at for s in shots if s.carry_distance]
                carries = [s.carry_distance for s in shots if s.carry_distance]
                if dates and carries:
                    ax.plot(dates, carries, color='#ff4d4d', linewidth=2, marker='o', markersize=4)
                    ax.set_xlabel('Date')
                    ax.set_ylabel('Carry Distance (yds)')
                    ax.set_title('Carry Distance Trend')
                    plt.xticks(rotation=45)
                    plt.tight_layout()
            
            elif "speed trends" in chart_type or "speed" in chart_type:
                # Line chart of club and ball speed
                dates = [s.recorded_at for s in shots if s.club_speed and s.ball_speed]
                club_speeds = [s.club_speed for s in shots if s.club_speed and s.ball_speed]
                ball_speeds = [s.ball_speed for s in shots if s.club_speed and s.ball_speed]
                if dates and club_speeds and ball_speeds:
                    ax.plot(dates, club_speeds, color='#ff4d4d', linewidth=2, marker='o', markersize=4, label='Club Speed')
                    ax.plot(dates, ball_speeds, color='#238636', linewidth=2, marker='s', markersize=4, label='Ball Speed')
                    ax.set_xlabel('Date')
                    ax.set_ylabel('Speed (mph)')
                    ax.set_title('Speed Trends')
                    ax.legend()
                    plt.xticks(rotation=45)
                    plt.tight_layout()
            
            elif "distance trends" in chart_type or "distance" in chart_type:
                # Line chart of carry and total distance
                dates = [s.recorded_at for s in shots if s.carry_distance and s.total_distance]
                carries = [s.carry_distance for s in shots if s.carry_distance and s.total_distance]
                totals = [s.total_distance for s in shots if s.carry_distance and s.total_distance]
                if dates and carries and totals:
                    ax.plot(dates, carries, color='#ff4d4d', linewidth=2, marker='o', markersize=4, label='Carry')
                    ax.plot(dates, totals, color='#238636', linewidth=2, marker='s', markersize=4, label='Total')
                    ax.set_xlabel('Date')
                    ax.set_ylabel('Distance (yds)')
                    ax.set_title('Distance Trends')
                    ax.legend()
                    plt.xticks(rotation=45)
                    plt.tight_layout()
            
            elif "spin rate trends" in chart_type or "spin" in chart_type:
                # Line chart of spin rate
                dates = [s.recorded_at for s in shots if s.spin_rate]
                spins = [s.spin_rate for s in shots if s.spin_rate]
                if dates and spins:
                    ax.plot(dates, spins, color='#ff4d4d', linewidth=2, marker='o', markersize=4)
                    ax.set_xlabel('Date')
                    ax.set_ylabel('Spin Rate (rpm)')
                    ax.set_title('Spin Rate Trend')
                    plt.xticks(rotation=45)
                    plt.tight_layout()
            
            elif "dispersion" in chart_type:
                # Scatter plot of carry distance vs launch direction
                carries = [s.carry_distance for s in shots if s.carry_distance and s.launch_direction]
                directions = [s.launch_direction for s in shots if s.carry_distance and s.launch_direction]
                if carries and directions:
                    ax.scatter(directions, carries, color='#ff4d4d', alpha=0.6, s=50)
                    ax.set_xlabel('Launch Direction (deg)')
                    ax.set_ylabel('Carry Distance (yds)')
                    ax.set_title('Dispersion Analysis')
                    plt.tight_layout()
            
            elif "club comparison" in chart_type or "club" in chart_type:
                # Bar chart comparing clubs
                from collections import defaultdict
                club_data = defaultdict(list)
                for shot in shots:
                    if shot.session_id:
                        for sess in sessions:
                            if sess.id == shot.session_id and sess.club and shot.carry_distance:
                                club_data[sess.club].append(shot.carry_distance)
                
                if club_data:
                    clubs = list(club_data.keys())
                    avg_carries = [sum(club_data[c]) / len(club_data[c]) for c in clubs]
                    ax.bar(clubs, avg_carries, color='#ff4d4d')
                    ax.set_xlabel('Club')
                    ax.set_ylabel('Avg Carry Distance (yds)')
                    ax.set_title('Club Comparison')
                    plt.xticks(rotation=45)
                    plt.tight_layout()
            
            elif "distribution" in chart_type:
                # Histogram of carry distance
                carries = [s.carry_distance for s in shots if s.carry_distance]
                if carries:
                    ax.hist(carries, bins=20, color='#ff4d4d', edgecolor='#30363d', alpha=0.7)
                    ax.set_xlabel('Carry Distance (yds)')
                    ax.set_ylabel('Frequency')
                    ax.set_title('Carry Distance Distribution')
                    plt.tight_layout()
            
            elif "box plot" in chart_type:
                # Box plot of carry distance
                carries = [s.carry_distance for s in shots if s.carry_distance]
                if carries:
                    ax.boxplot(carries, vert=True, patch_artist=True,
                              boxprops=dict(facecolor='#ff4d4d', alpha=0.7),
                              medianprops=dict(color='#c9d1d9', linewidth=2))
                    ax.set_ylabel('Carry Distance (yds)')
                    ax.set_title('Carry Distance Box Plot')
                    plt.tight_layout()
            
            else:
                # Default: simple line chart of carry distance
                dates = [s.recorded_at for s in shots if s.carry_distance]
                carries = [s.carry_distance for s in shots if s.carry_distance]
                if dates and carries:
                    ax.plot(dates, carries, color='#ff4d4d', linewidth=2, marker='o', markersize=4)
                    ax.set_xlabel('Date')
                    ax.set_ylabel('Carry Distance (yds)')
                    ax.set_title(section.title)
                    plt.xticks(rotation=45)
                    plt.tight_layout()
            
            fig.savefig(chart_path, dpi=150, facecolor='#0f1116', edgecolor='none', bbox_inches='tight')
            plt.close(fig)
            
            logger.info("Generated chart: %s", chart_path)
            return chart_path
            
        except Exception as e:
            logger.error("Error generating chart: %s", e, exc_info=True)
            return None
    
    def _generate_csv(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
    ) -> bool:
        """Generate CSV report."""
        try:
            import csv
            from sqlalchemy.orm import Session
            from core.session_manager import ShotModel, SessionModel
            
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            if not shots:
                logger.warning("No shots to export to CSV")
                return False
            
            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                
                # Write header
                headers = [
                    "Date", "Time", "Session", "Club", "Club Speed (mph)", "Ball Speed (mph)",
                    "Launch Angle (deg)", "Spin Rate (rpm)", "Carry Distance (yds)",
                    "Total Distance (yds)", "Side Spin (rpm)", "Back Spin (rpm)",
                    "Launch Direction (deg)", "Apex Height (yds)", "Descent Angle (deg)",
                    "Smash Factor", "Dynamic Loft (deg)", "Attack Angle (deg)",
                    "Club Path (deg)", "Face Angle (deg)", "Notes"
                ]
                writer.writerow(headers)
                
                # Get session map
                session_map = {}
                club_map = {}
                with Session(self.session_manager.engine) as session:
                    for shot in shots:
                        if shot.session_id not in session_map:
                            sess = session.get(SessionModel, shot.session_id)
                            session_map[shot.session_id] = sess.name if sess else "Unknown"
                            club_map[shot.session_id] = sess.club if sess and sess.club else ""
                
                # Write data rows
                for shot in shots:
                    row = [
                        shot.recorded_at.strftime("%Y-%m-%d"),
                        shot.recorded_at.strftime("%H:%M:%S"),
                        session_map.get(shot.session_id, "Unknown"),
                        club_map.get(shot.session_id, ""),
                        f"{shot.club_speed:.1f}" if shot.club_speed else "",
                        f"{shot.ball_speed:.1f}" if shot.ball_speed else "",
                        f"{shot.launch_angle:.1f}" if shot.launch_angle else "",
                        f"{shot.spin_rate:.0f}" if shot.spin_rate else "",
                        f"{shot.carry_distance:.1f}" if shot.carry_distance else "",
                        f"{shot.total_distance:.1f}" if shot.total_distance else "",
                        f"{shot.side_spin:.0f}" if shot.side_spin else "",
                        f"{shot.back_spin:.0f}" if shot.back_spin else "",
                        f"{shot.launch_direction:.1f}" if shot.launch_direction else "",
                        f"{shot.apex_height:.1f}" if shot.apex_height else "",
                        f"{shot.descent_angle:.1f}" if shot.descent_angle else "",
                        f"{shot.smash_factor:.2f}" if shot.smash_factor else "",
                        f"{shot.dynamic_loft:.1f}" if shot.dynamic_loft else "",
                        f"{shot.attack_angle:.1f}" if shot.attack_angle else "",
                        f"{shot.club_path:.1f}" if shot.club_path else "",
                        f"{shot.face_angle:.1f}" if shot.face_angle else "",
                        shot.notes or "",
                    ]
                    writer.writerow(row)
            
            logger.info("Generated CSV report: %s", output_path)
            return True
        except Exception as e:
            logger.error("Error generating CSV report: %s", e, exc_info=True)
            return False
    
    def _generate_excel(
        self,
        template: ReportTemplate,
        output_path: Path,
        session_ids: Optional[list[int]],
        shot_ids: Optional[list[int]],
        date_range: Optional[tuple[datetime, datetime]],
        club_filter: Optional[str],
    ) -> bool:
        """Generate Excel report."""
        try:
            import pandas as pd
            from sqlalchemy.orm import Session
            from core.session_manager import ShotModel, SessionModel
            
            shots, sessions = self._get_report_data(session_ids, shot_ids, date_range, club_filter)
            
            if not shots:
                logger.warning("No shots to export to Excel")
                return False
            
            # Get session map
            session_map = {}
            club_map = {}
            with Session(self.session_manager.engine) as session:
                for shot in shots:
                    if shot.session_id not in session_map:
                        sess = session.get(SessionModel, shot.session_id)
                        session_map[shot.session_id] = sess.name if sess else "Unknown"
                        club_map[shot.session_id] = sess.club if sess and sess.club else ""
            
            # Build data rows
            rows = []
            for shot in shots:
                row = {
                    "Date": shot.recorded_at.strftime("%Y-%m-%d"),
                    "Time": shot.recorded_at.strftime("%H:%M:%S"),
                    "Session": session_map.get(shot.session_id, "Unknown"),
                    "Club": club_map.get(shot.session_id, ""),
                    "Club Speed (mph)": shot.club_speed,
                    "Ball Speed (mph)": shot.ball_speed,
                    "Launch Angle (deg)": shot.launch_angle,
                    "Spin Rate (rpm)": shot.spin_rate,
                    "Carry Distance (yds)": shot.carry_distance,
                    "Total Distance (yds)": shot.total_distance,
                    "Side Spin (rpm)": shot.side_spin,
                    "Back Spin (rpm)": shot.back_spin,
                    "Launch Direction (deg)": shot.launch_direction,
                    "Apex Height (yds)": shot.apex_height,
                    "Descent Angle (deg)": shot.descent_angle,
                    "Smash Factor": shot.smash_factor,
                    "Dynamic Loft (deg)": shot.dynamic_loft,
                    "Attack Angle (deg)": shot.attack_angle,
                    "Club Path (deg)": shot.club_path,
                    "Face Angle (deg)": shot.face_angle,
                    "Notes": shot.notes or "",
                }
                rows.append(row)
            
            df = pd.DataFrame(rows)
            df.to_excel(output_path, index=False, engine='openpyxl')
            
            logger.info("Generated Excel report: %s", output_path)
            return True
        except ImportError:
            logger.error("pandas and openpyxl required for Excel export. Install with: pip install pandas openpyxl")
            return False
        except Exception as e:
            logger.error("Error generating Excel report: %s", e, exc_info=True)
            return False
