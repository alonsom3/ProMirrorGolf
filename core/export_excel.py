"""Excel export functionality with formatting."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def export_session_to_excel(session, shots: list, output_path: Path) -> bool:
    """Export a session and its shots to Excel with formatting.
    
    Args:
        session: SessionModel instance
        shots: List of ShotModel instances
        output_path: Path to output Excel file
        
    Returns:
        True if successful, False otherwise
    """
    try:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            return False
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Session Data"
        
        # Header style
        header_fill = PatternFill(start_color="FF4D4D", end_color="FF4D4D", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Session information section
        ws['A1'] = "Session Information"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:B1')
        
        row = 2
        session_data = [
            ("Name", session.name),
            ("Club", session.club or ""),
            ("Notes", session.notes or ""),
            ("Started", session.started_at.strftime("%Y-%m-%d %H:%M:%S")),
            ("Ended", session.ended_at.strftime("%Y-%m-%d %H:%M:%S") if session.ended_at else ""),
        ]
        
        for label, value in session_data:
            ws[f'A{row}'] = label
            ws[f'A{row}'].font = Font(bold=True)
            ws[f'B{row}'] = value
            row += 1
        
        row += 2
        
        # Shots header
        headers = [
            "Time", "Club Speed (mph)", "Ball Speed (mph)", "Launch Angle (deg)",
            "Spin Rate (rpm)", "Carry Distance (yds)", "Total Distance (yds)",
            "Side Spin (rpm)", "Back Spin (rpm)", "Launch Direction (deg)",
            "Apex Height (yds)", "Descent Angle (deg)", "Smash Factor",
            "Dynamic Loft (deg)", "Attack Angle (deg)", "Club Path (deg)",
            "Face Angle (deg)", "Notes"
        ]
        
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        row += 1
        
        # Shot data
        for shot in shots:
            data = [
                shot.recorded_at.strftime("%H:%M:%S"),
                shot.club_speed,
                shot.ball_speed,
                shot.launch_angle,
                shot.spin_rate,
                shot.carry_distance,
                shot.total_distance,
                shot.side_spin,
                shot.back_spin,
                shot.launch_direction,
                shot.apex_height,
                shot.descent_angle,
                shot.smash_factor,
                shot.dynamic_loft,
                shot.attack_angle,
                shot.club_path,
                shot.face_angle,
                shot.notes or "",
            ]
            
            for col, value in enumerate(data, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = border
                if isinstance(value, (int, float)):
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')
            
            row += 1
        
        # Auto-adjust column widths
        for col in range(1, len(headers) + 1):
            column_letter = get_column_letter(col)
            max_length = 0
            for cell in ws[column_letter]:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(output_path)
        logger.info("Exported session %d to Excel: %s", session.id, output_path)
        return True
    except Exception as e:
        logger.error("Error exporting to Excel: %s", e, exc_info=True)
        return False

